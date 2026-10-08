import re
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request as UrlRequest, urlopen

from fastapi import FastAPI, Request, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from database.db import (init_db, list_machines, get_machine, create_machine, set_active,
    update_machine_settings, recent_events, add_event,
    delete_machine as delete_machine_record,
)

from monitoring.adapter import MonitoringManager

BASE_DIR = Path(__file__).resolve().parent.parent
app = FastAPI(title="CNC Monitoring Dashboard")
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")
manager = MonitoringManager()

def machine_names_by_camera(machines):
    return {
        f"camera_{int(machine['id']):02d}": machine["name"]
        for machine in machines
    }

def add_machine_names_to_events(events, machine_names):
    for event in events:
        event["machine_name"] = machine_names.get(
            event.get("camera_id"),
            "Unregistered machine",
        )
    return events

def excel_text(value):
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text

def create_events_workbook(events):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Events"
    headers = [
        "Time",
        "Machine",
        "Event",
        "Video time (seconds)",
        "Track IDs",
        "Details",
        "Screenshot",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="top")

    for event in events:
        sheet.append([
            excel_text(event.get("timestamp", "")),
            excel_text(event.get("machine_name", "Unregistered machine")),
            excel_text(event.get("event_type", "")),
            event.get("video_time", 0),
            excel_text(event.get("track_ids", "")),
            excel_text(event.get("details", "")),
            excel_text(event.get("screenshot_url") or ""),
        ])

    widths = (24, 28, 32, 22, 18, 72, 64)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=index).column_letter
        ].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

def events_download_filename(machine_name=None):
    if not machine_name:
        return "cnc-events.xlsx"
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "-", machine_name).strip("-")
    return f"{safe_name or 'machine'}-events.xlsx"

def camera_host(rtsp_url: str):
    from urllib.parse import urlsplit
    parsed = urlsplit(rtsp_url.strip())
    if parsed.scheme not in ("rtsp", "rtsps") or not parsed.hostname:
        raise HTTPException(400, "Camera URL must contain a valid RTSP host")
    host = parsed.hostname.lower()
    if ":" in host and not host.startswith("["):
        host = f"[{host}]"
    if parsed.port:
        host = f"{host}:{parsed.port}"
    return host

def backend_stream(path: str):
    request = UrlRequest(f"{manager.backend_url}{quote(path, safe='/%')}")
    try:
        response = urlopen(request, timeout=10)
    except (HTTPError, URLError) as exc:
        raise HTTPException(502, f"CNC backend stream unavailable: {exc}") from exc
    try:
        while chunk := response.read(64 * 1024):
            yield chunk
    finally:
        response.close()

@app.on_event("startup")
def startup():
    init_db()

@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request):
    machines = []
    for machine in list_machines():
        machine["runtime"] = manager.status(machine["id"])
        try:
            machine["zones_saved"] = bool(manager.zones_status(machine["id"]).get("saved"))
        except RuntimeError:
            machine["zones_saved"] = False
        machines.append(machine)
    try:
        events = manager.events()
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend events unavailable: {exc}") from exc
    add_machine_names_to_events(events, machine_names_by_camera(machines))
    events_by_camera = {}
    for event in events:
        events_by_camera.setdefault(event["camera_id"], []).append(event)
    for machine in machines:
        machine["events"] = events_by_camera.get(
            f"camera_{int(machine['id']):02d}", []
        )
    return templates.TemplateResponse(
        request=request,
        name="dashboard.html",
        context={
            "machines": machines,
            "events": events,
        },
    )

@app.get("/machines/{machine_id}/monitor", response_class=HTMLResponse)
def machine_preview_page(request: Request, machine_id: int):
    machine = get_machine(machine_id)
    if not machine:
        raise HTTPException(404, "Machine not found")
    machine["runtime"] = manager.status(machine_id)
    try:
        machine["zones_saved"] = bool(manager.zones_status(machine_id).get("saved"))
    except RuntimeError:
        machine["zones_saved"] = False
    try:
        events = manager.events()
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend events unavailable: {exc}") from exc
    camera_id = f"camera_{int(machine['id']):02d}"
    machine_events = [
        event for event in events
        if event.get("camera_id") == camera_id
    ]
    add_machine_names_to_events(machine_events, machine_names_by_camera(list_machines()))
    return templates.TemplateResponse(
        request=request,
        name="machine_preview.html",
        context={"machine": machine, "events": machine_events},
    )

@app.post("/machines")
async def add_machine(request: Request, name: str = Form(...),
    rtsp_url: str = Form(...),
    max_persons: int = Form(...),
    zone_count: int = Form(1),
    multiple_limit_seconds: int = Form(120),
    absence_limit_seconds: int = Form(300),
):
    if not rtsp_url.lower().startswith(("rtsp://", "rtsps://")):
        raise HTTPException(400, "Camera URL must start with rtsp:// or rtsps://")
    if max_persons < 1 or max_persons > 50:
        raise HTTPException(400, "Maximum people must be between 1 and 50")
    if not 1 <= zone_count <= 20:
        raise HTTPException(400, "Number of zones must be between 1 and 20")
    form = await request.form()
    zone_limits = {}
    for index in range(1, zone_count + 1):
        try:
            zone_limits[f"zone_{index}"] = max(0, min(50, int(
                form.get(f"zone_limit_{index}", max_persons)
            )))
        except (TypeError, ValueError) as exc:
            raise HTTPException(400, f"Invalid limit for zone {index}") from exc
    machine_id = create_machine(
        name.strip(), rtsp_url.strip(), max_persons, multiple_limit_seconds,
        absence_limit_seconds, zone_limits,
    )
    return RedirectResponse(url=f"/", status_code=303)

@app.post("/machines/{machine_id}/start")
def start_machine(machine_id: int):
    machine = get_machine(machine_id)
    if not machine:
        raise HTTPException(404, "Machine not found")
    try:
        raise HTTPException(400, "Use the dashboard zone editor before starting monitoring")
    except RuntimeError as exc:
        raise HTTPException(502, str(exc)) from exc
    set_active(machine_id, True)
    return RedirectResponse(url="/", status_code=303)

@app.post("/machines/{machine_id}/start-monitoring")
async def start_monitoring(machine_id: int, request: Request):
    machine = get_machine(machine_id)
    if not machine:
        raise HTTPException(404, "Machine not found")
    payload = await request.json()
    zones = payload.get("zones")
    if not isinstance(zones, list):
        raise HTTPException(400, "Zones must be an array")
    if not zones:
        try:
            manager.start(machine)
        except RuntimeError as exc:
            raise HTTPException(502, str(exc)) from exc
        set_active(machine_id, True)
        return {"started": True}
    if not 1 <= len(zones) <= 20:
        raise HTTPException(400, "Number of zones must be between 1 and 20")
    zone_limits = machine.get("zone_limits", {})
    zone_limits = {
        f"zone_{index}": max(0, min(50, int(
            zone_limits.get(f"zone_{index}", machine["max_persons"])
        )))
        for index in range(1, len(zones) + 1)
    }
    try:
        update_machine_settings(
            machine_id,
            machine["max_persons"],
            machine["multiple_limit_seconds"],
            machine["absence_limit_seconds"],
            zone_limits,
        )
        machine["zone_limits"] = zone_limits
        manager.start_with_zones(machine, zones)
    except RuntimeError as exc:
        raise HTTPException(502, str(exc)) from exc
    set_active(machine_id, True)
    return {"started": True}

@app.get("/machines/{machine_id}/preview")
def machine_preview(machine_id: int):
    machine = get_machine(machine_id)
    if not machine:
        raise HTTPException(404, "Machine not found")
    try:
        preview = manager.preview(machine)
    except RuntimeError as exc:
        raise HTTPException(502, str(exc)) from exc
    return StreamingResponse(iter([preview]), media_type="image/jpeg")

@app.get("/machines/{machine_id}/zones")
def machine_zones(machine_id: int):
    if not get_machine(machine_id):
        raise HTTPException(404, "Machine not found")
    try:
        return manager.zones_status(machine_id)
    except RuntimeError as exc:
        raise HTTPException(
            502, f"CNC backend saved zones unavailable: {exc}"
        ) from exc

@app.post("/machines/{machine_id}/zones")
async def save_machine_zones(machine_id: int, request: Request):
    machine = get_machine(machine_id)
    if not machine:
        raise HTTPException(404, "Machine not found")

    payload = await request.json()
    zones = payload.get("zones") if isinstance(payload, dict) else None
    if not isinstance(zones, list) or not 1 <= len(zones) <= 20:
        raise HTTPException(400, "Number of zones must be between 1 and 20")
    try:
        manager.save_zones(machine_id, zones)
        runtime = manager.status(machine_id)
        zone_limits = machine.get("zone_limits", {})
        zone_limits = {
            f"zone_{index}": max(0, min(50, int(
                zone_limits.get(f"zone_{index}", machine["max_persons"])
            )))
            for index in range(1, len(zones) + 1)
        }
        update_machine_settings(
            machine_id,
            machine["max_persons"],
            machine["multiple_limit_seconds"],
            machine["absence_limit_seconds"],
            zone_limits,
        )
        machine["zone_limits"] = zone_limits
        if runtime.get("running"):
            manager.start_with_zones(machine, zones)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(502, f"Could not save or apply work zones: {exc}") from exc

    return {"saved": True, "count": len(zones), "running": bool(runtime.get("running"))}

@app.post("/machines/{machine_id}/settings")
async def update_settings( request: Request,
    machine_id: int,
    max_persons: int = Form(...),
    multiple_limit_seconds: int = Form(...),
    absence_limit_seconds: int = Form(...),
):
    machine = get_machine(machine_id)
    if not machine:
        raise HTTPException(404, "Machine not found")
    if not 1 <= max_persons <= 50:
        raise HTTPException(400, "Maximum people must be between 1 and 50")
    if multiple_limit_seconds < 1 or absence_limit_seconds < 1:
        raise HTTPException(400, "Delays must be at least 1 second")
    form = await request.form()
    zone_limits = {}
    for key, value in form.items():
        if key.startswith("zone_limit_"):
            zone_number = key.removeprefix("zone_limit_")
            try:
                zone_limits[f"zone_{int(zone_number)}"] = max(0, min(50, int(value)))
            except ValueError as exc:
                raise HTTPException(400, f"Invalid limit for zone {zone_number}") from exc
    if not zone_limits:
        zone_limits = {"zone_1": max_persons}
    update_machine_settings(machine_id,
        max_persons,
        multiple_limit_seconds,
        absence_limit_seconds,
        zone_limits,
    )
    
    if manager.status(machine_id)["running"]:
        manager.stop(machine_id)
        manager.start(get_machine(machine_id))
    return RedirectResponse(url="/", status_code=303)

@app.post("/machines/{machine_id}/stop")
def stop_machine(machine_id: int):
    if not get_machine(machine_id):
        raise HTTPException(404, "Machine not found")
    manager.stop(machine_id)
    set_active(machine_id, False)
    return RedirectResponse(url="/", status_code=303)

@app.post("/machines/{machine_id}/delete")
def delete_machine(machine_id: int):
    if not get_machine(machine_id):
        raise HTTPException(404, "Machine not found")
    try:
        manager.delete_machine(machine_id)
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend could not delete machine: {exc}") from exc
    if not delete_machine_record(machine_id):
        raise HTTPException(404, "Machine not found")
    return RedirectResponse(url="/", status_code=303)

@app.get("/api/machines/{machine_id}/status")
def machine_status(machine_id: int):
    if not get_machine(machine_id):
        raise HTTPException(404, "Machine not found")
    return manager.status(machine_id)

@app.post("/api/camera-credentials")
def save_credentials(rtsp_url: str = Form(...), username: str = Form(...), password: str = Form(...)):
    from database.db import save_camera_credentials
    host = camera_host(rtsp_url)
    username = username.strip()
    if not username or not password:
        raise HTTPException(400, "Camera username and password are required")
    save_camera_credentials(host, username, password)
    return {"saved": True}

@app.get("/backend/detection/{camera_id}/stream")
def machine_stream(camera_id: str):
    return StreamingResponse(
        backend_stream(f"/detection/{camera_id}/stream"),
        media_type="multipart/x-mixed-replace; boundary=frame",
    )

@app.get("/backend/outputs/{file_path:path}")
def output_file(file_path: str):
    return StreamingResponse(
        backend_stream(f"/outputs/{file_path}"),
        media_type="image/jpeg",
    )

@app.get("/api/events")
def events():
    try:
        return manager.events()
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend events unavailable: {exc}") from exc

@app.get("/api/events/export.xlsx")
def export_events(machine_id: int | None = None):
    machine = get_machine(machine_id) if machine_id is not None else None
    if machine_id is not None and machine is None:
        raise HTTPException(404, "Machine not found")
    try:
        records = manager.events(limit=10000)
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend events unavailable: {exc}") from exc

    names = machine_names_by_camera(list_machines())
    add_machine_names_to_events(records, names)
    if machine is not None:
        camera_id = f"camera_{int(machine['id']):02d}"
        records = [
            event for event in records
            if event.get("camera_id") == camera_id
        ]
    if not records:
        raise HTTPException(404, "No events available to export")

    workbook_file = create_events_workbook(records)
    filename = events_download_filename(
        machine["name"] if machine is not None else None
    )
    return StreamingResponse(
        workbook_file,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f"attachment; filename*=UTF-8''{quote(filename, safe='')}"
        },
    )

@app.post("/api/events/delete")
def delete_events(event_ids: list[str] = Form(...)):
    try:
        manager.delete_events(event_ids)
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend event deletion failed: {exc}") from exc
    return RedirectResponse(url="/", status_code=303)

@app.post("/api/events/{event_id}/delete")
def delete_event(event_id: str):
    try:
        manager.delete_event(event_id)
    except RuntimeError as exc:
        raise HTTPException(502, f"CNC backend event deletion failed: {exc}") from exc
    return RedirectResponse(url="/", status_code=303)
