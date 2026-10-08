import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "cnc_frontend"))

from web.app import (
    add_machine_names_to_events,
    create_events_workbook,
    excel_text,
    events_download_filename,
    export_events,
    machine_names_by_camera,
)


class EventExportTests(unittest.TestCase):
    def test_event_machine_names_use_registered_machine_names(self):
        names = machine_names_by_camera([
            {"id": 3, "name": "Press Line"},
            {"id": 12, "name": "Mill 12"},
        ])
        events = [
            {"camera_id": "camera_03"},
            {"camera_id": "camera_12"},
            {"camera_id": "camera_99"},
        ]

        add_machine_names_to_events(events, names)

        self.assertEqual(
            [event["machine_name"] for event in events],
            ["Press Line", "Mill 12", "Unregistered machine"],
        )

    def test_excel_workbook_contains_machine_name_and_event_data(self):
        workbook_file = create_events_workbook([{
            "timestamp": "2026-10-08 03:30:00 PM",
            "machine_name": "Press Line",
            "event_type": "HEAD_DOWN_PHONE_VIOLATION",
            "video_time": 301.5,
            "track_ids": "7",
            "details": "Phone and head down for five minutes",
            "screenshot_url": "/outputs/cameras/camera_03/shot.jpg",
        }])

        workbook = load_workbook(workbook_file, read_only=True)
        rows = list(workbook["Events"].iter_rows(values_only=True))

        self.assertEqual(rows[0], (
            "Time",
            "Machine",
            "Event",
            "Video time (seconds)",
            "Track IDs",
            "Details",
            "Screenshot",
        ))
        self.assertEqual(rows[1][1], "Press Line")
        self.assertEqual(rows[1][2], "HEAD_DOWN_PHONE_VIOLATION")
        self.assertEqual(rows[1][3], 301.5)

    def test_export_file_names_are_safe_and_machine_specific(self):
        self.assertEqual(events_download_filename(), "cnc-events.xlsx")
        self.assertEqual(
            events_download_filename("CNC Press #1"),
            "CNC-Press-1-events.xlsx",
        )

    def test_excel_text_keeps_user_content_as_plain_text(self):
        self.assertEqual(excel_text("=HYPERLINK(\"url\")"), "'=HYPERLINK(\"url\")")

    def test_export_rejects_empty_event_list(self):
        with patch("web.app.manager.events", return_value=[]):
            with self.assertRaisesRegex(HTTPException, "No events available to export") as error:
                export_events()

        self.assertEqual(error.exception.status_code, 404)

    def test_machine_export_rejects_when_machine_has_no_events(self):
        events = [{"camera_id": "camera_02", "event_type": "VIOLATION"}]
        machine = {"id": 3, "name": "Press Line"}
        with (
            patch("web.app.manager.events", return_value=events),
            patch("web.app.get_machine", return_value=machine),
            patch("web.app.list_machines", return_value=[machine]),
        ):
            with self.assertRaisesRegex(HTTPException, "No events available to export") as error:
                export_events(machine_id=3)

        self.assertEqual(error.exception.status_code, 404)


if __name__ == "__main__":
    unittest.main()
